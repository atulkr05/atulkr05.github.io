"""Rebuild assets/data/faceprint-atlas.xlsx from assets/data/atlas-data.js.

Usage (from the repository root):
    pip install openpyxl
    python3 _tools/atlas/build_xlsx.py

The workbook's Summary and index tabs use formulas. Excel, LibreOffice and Google Sheets
calculate them when the file is opened.
"""
import json
import os
import sys
from urllib.parse import quote

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit('openpyxl is required: pip install openpyxl')

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
SRC = os.path.join(REPO, 'assets', 'data', 'atlas-data.js')
OUT = os.path.join(REPO, 'assets', 'data', 'faceprint-atlas.xlsx')


def load_data(path):
    text = open(path, encoding='utf-8').read()
    start, end = text.index('window.ATLAS'), text.rindex('}')
    body = text[text.index('{', start):end + 1]
    try:
        return json.loads(body)
    except json.JSONDecodeError as err:
        lines = body.splitlines()
        context = lines[err.lineno - 1] if 0 < err.lineno <= len(lines) else ''
        sys.exit(f'atlas-data.js is not valid JSON near data line {err.lineno}: {err.msg}\n  {context.strip()}')


def validate(data):
    fams = {f['key'] for f in data['families']}
    groups = {g['key'] for g in data['groups']}
    ids, problems = set(), []
    required = ['id', 'title', 'authors', 'year', 'venue', 'venueType', 'family', 'tags', 'approach', 'findings',
                'datasets', 'models', 'threatModel', 'utility', 'limitations', 'code', 'github', 'arxiv', 'doi', 'url']
    for f in data['families']:
        if f['group'] not in groups:
            problems.append(f'family {f["key"]}: unknown group {f["group"]}')
    for p in data['papers']:
        missing = [k for k in required if k not in p]
        if missing:
            problems.append(f'{p.get("id", "?")}: missing {", ".join(missing)}')
        if p.get('id') in ids:
            problems.append(f'{p["id"]}: duplicate id')
        ids.add(p.get('id'))
        if p.get('family') not in fams:
            problems.append(f'{p.get("id")}: unknown family {p.get("family")}')
        if not isinstance(p.get('year'), int):
            problems.append(f'{p.get("id")}: year must be a number')
        if p.get('code') and not p.get('github'):
            problems.append(f'{p.get("id")}: code is true but github is empty')
    for d in data.get('duels', []):
        if d['attack'] not in ids:
            problems.append(f'duel attack {d["attack"]}: no such paper id')
    for pth in data.get('paths', []):
        for i in pth['ids']:
            if i not in ids:
                problems.append(f'reading path "{pth["name"]}": no such paper id {i}')
    if problems:
        sys.exit('atlas-data.js has problems:\n  ' + '\n  '.join(problems))


def paper_link(p):
    if p['arxiv']:
        return 'https://arxiv.org/abs/' + p['arxiv']
    if p['doi']:
        return 'https://doi.org/' + p['doi']
    return p['url'] or scholar(p)


def scholar(p):
    return 'https://scholar.google.com/scholar?q=' + quote('"' + p['title'] + '"')


def build(data):
    fam_by = {f['key']: f for f in data['families']}
    grp_by = {g['key']: g for g in data['groups']}
    papers = sorted(data['papers'], key=lambda p: (-p['year'], p['title']))
    for p in papers:
        p['group'] = fam_by[p['family']]['group']

    FONT = 'Arial'
    head_fill = PatternFill('solid', fgColor='0F172A')
    head_font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
    body = Font(name=FONT, size=9)
    link_font = Font(name=FONT, size=9, color='2563EB', underline='single')
    border = Border(bottom=Side(style='thin', color='DDDDDD'))
    tint = {'g1': 'DCE9F8', 'g2': 'FBE2D6', 'g3': 'D3F1E5', 'g4': 'FBEFCC', 'g5': 'FADDE8', 'g6': 'D6EDD6', 'g7': 'E0DDF3', 'g8': 'F9DADA'}
    cols = [('#', 5), ('ID', 16), ('Paper title', 48), ('Authors', 36), ('Year', 7), ('Venue', 26), ('Venue type', 11),
            ('Family', 26), ('Chart group', 22), ('Datasets', 34), ('Models / backbones', 34), ('Approach', 60),
            ('Key findings', 60), ('Threat model / protects against', 36), ('Utility preserved', 28),
            ('Limitations & known attacks', 50), ('Tags', 28), ('Code available', 10), ('GitHub / code link', 40),
            ('Paper link', 40), ('Google Scholar', 30)]

    def table(ws, rows, title, note):
        ws.cell(row=1, column=1, value=title).font = Font(name=FONT, bold=True, size=14)
        ws.cell(row=2, column=1, value=note).font = Font(name=FONT, italic=True, size=9, color='555555')
        hdr = 4
        for j, (name, width) in enumerate(cols, 1):
            c = ws.cell(row=hdr, column=j, value=name)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical='center', wrap_text=True)
            ws.column_dimensions[get_column_letter(j)].width = width
        ws.row_dimensions[hdr].height = 30
        for i, p in enumerate(rows, 1):
            r = hdr + i
            values = [i, p['id'], p['title'], p['authors'], p['year'], p['venue'], p['venueType'], fam_by[p['family']]['name'],
                      grp_by[p['group']]['name'], p['datasets'], p['models'], p['approach'], p['findings'], p['threatModel'],
                      p['utility'], p['limitations'], ', '.join(p['tags']), 'Yes' if p['code'] else 'No', p['github'],
                      paper_link(p), 'Search']
            for j, v in enumerate(values, 1):
                c = ws.cell(row=r, column=j, value=v)
                c.font, c.border = body, border
                c.alignment = Alignment(vertical='top', wrap_text=j not in (1, 2, 5, 7, 18, 19, 20, 21))
            ws.cell(row=r, column=8).fill = PatternFill('solid', fgColor=tint.get(p['group'], 'EEEEEE'))
            if p['github']:
                c = ws.cell(row=r, column=19); c.hyperlink = p['github']; c.font = link_font
            c = ws.cell(row=r, column=20); c.hyperlink = paper_link(p); c.font = link_font
            c = ws.cell(row=r, column=21); c.hyperlink = scholar(p); c.font = link_font
            ws.row_dimensions[r].height = 96
        ws.freeze_panes = ws.cell(row=hdr + 1, column=4)
        if rows:
            ws.auto_filter.ref = f'A{hdr}:{get_column_letter(len(cols))}{hdr + len(rows)}'
        return hdr

    wb = Workbook()
    ws = wb.active
    ws.title = 'README'
    years = [p['year'] for p in papers]
    readme = [
        ('Faceprint Atlas: facial privacy protection literature workbook', Font(name=FONT, bold=True, size=16)),
        (f'{len(papers)} entries from {min(years)} to {max(years)}, densest coverage from 2018. Last updated {data.get("updated", "")}. '
         'Interactive version: https://atulkr05.github.io/atlas.html', Font(name=FONT, size=10)),
        ('', None),
        ('How to use this workbook', Font(name=FONT, bold=True, size=12)),
        ('"All papers" holds every entry with the full record. Year tabs (Before 2018, then 2018 onward) and family tabs hold the same '
         'columns filtered. "Summary" counts entries by year and family with live COUNTIFS formulas over "All papers". '
         '"Attacks vs defenses" pairs defense papers with the papers that tested them. "Datasets index" and "Models index" count how '
         'often common datasets and backbones appear.', Font(name=FONT, size=10)),
        ('', None),
        ('Families', Font(name=FONT, bold=True, size=12)),
    ]
    for text, font in readme:
        ws.append([text])
        if font:
            ws.cell(row=ws.max_row, column=1).font = font
    for f in data['families']:
        ws.append([f['name'], f'{f["description"]} Examples: {f["examples"]}.'])
        ws.cell(row=ws.max_row, column=1).font = Font(name=FONT, bold=True, size=10)
        ws.cell(row=ws.max_row, column=1).fill = PatternFill('solid', fgColor=tint.get(f['group'], 'EEEEEE'))
        ws.cell(row=ws.max_row, column=2).font = Font(name=FONT, size=10)
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical='top')
    ws.append([])
    ws.append(['Caveats'])
    ws.cell(row=ws.max_row, column=1).font = Font(name=FONT, bold=True, size=12)
    ws.append(['', 'Curated catalogue, not a systematic review. Venues and years were checked against arXiv metadata and proceedings; '
                   'preprints may since have been published. Findings are summaries, so read the paper before citing a number. '
                   'Code links point only to repositories that could be confirmed.'])
    ws.cell(row=ws.max_row, column=2).font = Font(name=FONT, size=10)
    ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 48
    ws.column_dimensions['B'].width = 120
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                c.font = Font(name=FONT, size=10)

    ws = wb.create_sheet('All papers')
    hdr = table(ws, papers, 'All papers', 'Every entry, newest first. Use the filter arrows on the header row.')
    first, last = hdr + 1, hdr + len(papers)
    rng = lambda col: f"'All papers'!${col}${first}:${col}${last}"
    year_rng, fam_rng, vt_rng, ds_rng, md_rng, code_rng = rng('E'), rng('H'), rng('G'), rng('J'), rng('K'), rng('R')

    ws = wb.create_sheet('Summary')
    ws['A1'] = 'Entries by year and family'
    ws['A1'].font = Font(name=FONT, bold=True, size=14)
    ws['A2'] = 'Live COUNTIFS over the "All papers" tab.'
    ws['A2'].font = Font(name=FONT, italic=True, size=9, color='555555')
    labels = ['Before 2018'] + [str(y) for y in range(2018, max(years) + 1)]
    for j, label in enumerate(['Family'] + labels + ['Total'], 1):
        c = ws.cell(row=4, column=j, value=label)
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal='center')
    for i, f in enumerate(data['families'], 5):
        ws.cell(row=i, column=1, value=f['name']).font = body
        ws.cell(row=i, column=1).fill = PatternFill('solid', fgColor=tint.get(f['group'], 'EEEEEE'))
        for j, label in enumerate(labels, 2):
            cond = '"<2018"' if label == 'Before 2018' else label
            c = ws.cell(row=i, column=j, value=f'=COUNTIFS({fam_rng},$A{i},{year_rng},{cond})')
            c.font, c.alignment = body, Alignment(horizontal='center')
        last_col = get_column_letter(len(labels) + 1)
        ws.cell(row=i, column=len(labels) + 2, value=f'=SUM(B{i}:{last_col}{i})').font = Font(name=FONT, bold=True, size=9)
    total_row = 5 + len(data['families'])
    ws.cell(row=total_row, column=1, value='Total').font = Font(name=FONT, bold=True, size=9)
    for j in range(2, len(labels) + 3):
        col = get_column_letter(j)
        c = ws.cell(row=total_row, column=j, value=f'=SUM({col}5:{col}{total_row - 1})')
        c.font, c.alignment = Font(name=FONT, bold=True, size=9), Alignment(horizontal='center')
    ws.column_dimensions['A'].width = 48
    for j in range(2, len(labels) + 3):
        ws.column_dimensions[get_column_letter(j)].width = 11
    r = total_row + 3
    ws.cell(row=r, column=1, value='Other counts').font = Font(name=FONT, bold=True, size=12)
    extras = [('Entries with public code', f'=COUNTIF({code_rng},"Yes")'), ('Entries from 2018 onward', f'=COUNTIF({year_rng},">=2018")'),
              ('Preprints or under review', f'=COUNTIF({vt_rng},"Preprint")'), ('Grey literature (Other)', f'=COUNTIF({vt_rng},"Other")'),
              ('All entries', f'=COUNTA({fam_rng})')]
    for k, (label, formula) in enumerate(extras, r + 1):
        ws.cell(row=k, column=1, value=label).font = body
        ws.cell(row=k, column=2, value=formula).font = body
    ws.freeze_panes = 'B5'

    pre = [p for p in papers if p['year'] < 2018]
    table(wb.create_sheet('Before 2018'), pre, 'Before 2018', f'{len(pre)} entries before 2018: the foundations.')
    for y in range(2018, max(years) + 1):
        rows = [p for p in papers if p['year'] == y]
        table(wb.create_sheet(str(y)), rows, str(y), f'{len(rows)} entries published or first released in {y}.')

    tab_names = {'tpl': 'Identity protection', 'soft': 'Soft-biometric privacy', 'gen': 'De-identification', 'adv': 'Adversarial privacy',
                 'unl': 'Unlearnable images', 'mul': 'Model unlearning', 'dp': 'Differential privacy', 'obf': 'Blur & pixelation',
                 'atk': 'Attacks & evaluations', 'data': 'Training data & synthetic', 'surv': 'Surveys & society'}
    for key in ['tpl', 'soft', 'gen', 'adv', 'unl', 'mul', 'dp', 'obf', 'atk', 'data', 'surv'] + [f['key'] for f in data['families'] if f['key'] not in tab_names]:
        f = fam_by[key]
        rows = [p for p in papers if p['family'] == key]
        table(wb.create_sheet(tab_names.get(key, f['short'])[:31]), rows, f['name'], f'{f["description"]} {len(rows)} entries.')

    ws = wb.create_sheet('Attacks vs defenses')
    by_id = {p['id']: p for p in papers}
    heads = ['Defense (paper or family)', 'Defense year', 'Defense venue', 'Attack / evaluation paper', 'Attack year', 'Attack venue', 'What it showed', 'Attack paper link']
    for j, name in enumerate(heads, 1):
        c = ws.cell(row=1, column=j, value=name)
        c.font, c.fill = head_font, head_fill
    for i, d in enumerate(data.get('duels', []), 2):
        dp, ap = by_id.get(d['defense']), by_id[d['attack']]
        values = [dp['title'] if dp else d['defense'], dp['year'] if dp else '', dp['venue'] if dp else '', ap['title'], ap['year'], ap['venue'], d['verdict'], paper_link(ap)]
        for j, v in enumerate(values, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = body, border
            c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(row=i, column=8).hyperlink = paper_link(ap)
        ws.cell(row=i, column=8).font = link_font
    for j, width in enumerate([46, 9, 22, 46, 9, 22, 60, 36], 1):
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = 'A2'

    def index_sheet(title, rng_ref, items, note):
        sh = wb.create_sheet(title)
        sh['A1'] = title
        sh['A1'].font = Font(name=FONT, bold=True, size=14)
        sh['A2'] = note
        sh['A2'].font = Font(name=FONT, italic=True, size=9, color='555555')
        for j, name in enumerate(['Name', 'Entries mentioning it', 'Share of all entries'], 1):
            c = sh.cell(row=4, column=j, value=name)
            c.font, c.fill = head_font, head_fill
        for i, name in enumerate(items, 5):
            sh.cell(row=i, column=1, value=name).font = body
            sh.cell(row=i, column=2, value=f'=COUNTIF({rng_ref},"*{name}*")').font = body
            c = sh.cell(row=i, column=3, value=f'=IF(COUNTA({fam_rng})=0,0,B{i}/COUNTA({fam_rng}))')
            c.font, c.number_format = body, '0.0%'
        sh.column_dimensions['A'].width = 30
        sh.column_dimensions['B'].width = 20
        sh.column_dimensions['C'].width = 18
        sh.freeze_panes = 'A5'

    index_sheet('Datasets index', ds_rng, ['LFW', 'CelebA', 'CelebA-HQ', 'FFHQ', 'VGGFace2', 'CASIA-WebFace', 'MS1M', 'WebFace', 'FaceScrub', 'PubFig',
                                           'PIPA', 'FERET', 'FRGCv2', 'MOBIO', 'AgeDB', 'CFP-FP', 'IJB', 'MORPH', 'MUCT', 'Adience', 'CIFAR', 'ImageNet',
                                           'MNIST', 'Multi-PIE', 'Cityscapes', 'COCO', 'WikiArt', 'LADN', 'Lacuna'],
                'Wildcard COUNTIF over the Datasets column of "All papers". One entry can mention several datasets.')
    index_sheet('Models index', md_rng, ['ArcFace', 'FaceNet', 'CosFace', 'SphereFace', 'MobileFace', 'IR152', 'IRSE50', 'VGG-Face', 'AdaFace', 'MagFace',
                                         'ElasticFace', 'ResNet', 'ViT', 'StyleGAN', 'Stable Diffusion', 'DreamBooth', 'Diffusion', 'GAN', 'Eigenface', 'Face++',
                                         'Azure', 'Rekognition', 'Aliyun', 'CLIP', 'homomorphic', 'FHE', 'Paillier', 'Viola-Jones'],
                'Wildcard COUNTIF over the Models / backbones column of "All papers". One entry can mention several models.')
    return wb


if __name__ == '__main__':
    data = load_data(SRC)
    validate(data)
    wb = build(data)
    wb.save(OUT)
    print(f'Wrote {os.path.relpath(OUT, REPO)}: {len(data["papers"])} papers, {len(wb.sheetnames)} tabs.')
